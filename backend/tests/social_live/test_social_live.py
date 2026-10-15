"""Social Live ingestion + Medical 360 queries. Portable SQL, so these run on in-memory SQLite."""
import io
from datetime import datetime, timedelta, timezone

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select, update

from livepulse.modules.social_live import models as m
from livepulse.modules.social_live.retention import purge_expired

INGEST = {'Authorization': 'Bearer ingest-secret'}
READ = {'Authorization': 'Bearer read-secret'}
EXPORT = {'Authorization': 'Bearer export-secret'}
API = '/api/social-live'
MKT = '/api/marketing/tiktok'


async def seed(client, live_id='LIVE-1025', account='medical', tiktoker='lucia'):
    assert (await client.post(f'{API}/accounts', headers=INGEST, json={
        'id': account, 'name': 'Medical', 'username': '@medical.cirugias'})).status_code == 200
    assert (await client.post(f'{API}/tiktokers', headers=INGEST, json={
        'id': tiktoker, 'name': 'Lucía Ramos'})).status_code == 200
    started = (datetime.now(timezone.utc) - timedelta(hours=2)).isoformat()
    created = await client.post(f'{API}/lives', headers=INGEST, json={
        'live_id': live_id, 'tiktoker_id': tiktoker, 'account_id': account,
        'started_at': started})
    assert created.status_code == 201, created.text
    return live_id


async def test_scopes_are_separate(env):
    client, _ = env
    assert (await client.get(f'{MKT}/overview')).status_code == 401
    assert (await client.get(f'{MKT}/overview', headers={'Authorization': 'Bearer nope'})).status_code == 401
    assert (await client.get(f'{MKT}/overview', headers=INGEST)).status_code == 403
    assert (await client.post(f'{API}/tiktokers', headers=READ,
                              json={'id': 'x', 'name': 'X'})).status_code == 403
    assert (await client.get(f'{MKT}/overview', headers=READ)).status_code == 200
    await seed(client)
    assert (await client.get(f'{MKT}/lives/LIVE-1025/comments/export', headers=READ)).status_code == 403
    assert (await client.get(f'{MKT}/lives/LIVE-1025/comments/export', headers=EXPORT)).status_code == 200


async def test_create_live_is_idempotent_and_validates_references(env):
    client, _ = env
    await seed(client)
    again = await client.post(f'{API}/lives', headers=INGEST, json={
        'live_id': 'LIVE-1025', 'tiktoker_id': 'lucia', 'account_id': 'medical'})
    assert again.status_code == 200 and again.json()['created'] is False
    clash = await client.post(f'{API}/lives', headers=INGEST, json={
        'live_id': 'LIVE-1025', 'tiktoker_id': 'other', 'account_id': 'medical'})
    assert clash.status_code == 409
    missing = await client.post(f'{API}/lives', headers=INGEST, json={
        'tiktoker_id': 'ghost', 'account_id': 'medical'})
    assert missing.status_code == 422


async def test_events_consolidate_metrics_and_deduplicate_retries(env):
    client, _ = env
    live = await seed(client)
    batch = {'events': [
        {'event_type': 'comment', 'username': 'usuario123', 'event_id': 'c1',
         'text': 'quiero información 987654321'},
        {'event_type': 'comment', 'username': 'usuario123', 'event_id': 'c2',
         'text': 'me interesa el precio, 985//857#968'},
        {'event_type': 'comment', 'username': 'ana', 'event_id': 'c3', 'text': 'qué bonito'},
        {'event_type': 'like', 'value': 40, 'event_id': 'l1'},
        {'event_type': 'view', 'value': 300, 'event_id': 'v1'},
        {'event_type': 'viewer', 'value': 55}, {'event_type': 'viewer', 'value': 31},
        {'event_type': 'share', 'value': 3}, {'event_type': 'follow', 'value': 2},
        {'event_type': 'donation', 'value': 1}, {'event_type': 'direct_message'},
        {'event_type': 'phone_manual', 'phone': '+51 912 345 678', 'username': 'maria'},
    ]}
    first = (await client.post(f'{API}/lives/{live}/events', headers=INGEST, json=batch)).json()
    assert first['received'] == 12 and first['accepted'] == 12 and first['duplicates'] == 0
    retry = (await client.post(f'{API}/lives/{live}/events', headers=INGEST, json=batch)).json()
    assert retry['duplicates'] == 5  # only events that carried an event_id are idempotent

    metrics = (await client.get(f'{MKT}/lives/{live}/metrics', headers=READ)).json()
    assert metrics['comments_count'] == 3
    assert metrics['likes_count'] == 40 and metrics['views'] == 300
    assert metrics['peak_viewers'] == 55  # highest concurrent sample, never lowered
    assert metrics['average_viewers'] == 43.0  # (55 + 31) / 2 samples
    assert metrics['unique_viewers'] is None  # no source: not available
    assert metrics['interested_viewers'] == 1  # the same user twice counts once
    assert metrics['phones_captured'] == 3  # 987654321, 985857968, 912345678
    # Counter events without event_id are not idempotent: the retry added the share again.
    assert metrics['shares_count'] == 6
    assert metrics['leads_per_hour'] == pytest.approx(1.5, abs=0.1)  # 3 phones / 2 h


async def test_single_event_and_comment_shortcut_and_invalid_input(env):
    client, _ = env
    live = await seed(client)
    ok = await client.post(f'{API}/lives/{live}/events', headers=INGEST, json={
        'live_id': live, 'event_type': 'comment', 'username': 'usuario123',
        'text': 'quiero información 987654321', 'timestamp': '2026-10-01T19:23:10'})
    assert ok.status_code == 200, ok.text
    shortcut = await client.post(f'{API}/lives/{live}/comments', headers=INGEST,
                                 json={'username': 'ana', 'text': '9*8*1*545*41*1'})
    assert shortcut.status_code == 200
    assert (await client.post(f'{API}/lives/{live}/events', headers=INGEST,
                              json={'event_type': 'bogus'})).status_code == 422
    assert (await client.post(f'{API}/lives/{live}/events', headers=INGEST,
                              json={'event_type': 'comment', 'username': 'x'})).status_code == 422
    assert (await client.post(f'{API}/lives/NOPE/events', headers=INGEST,
                              json={'event_type': 'like'})).status_code == 404
    future = (datetime.now(timezone.utc) + timedelta(days=2)).isoformat()
    assert (await client.post(f'{API}/lives/{live}/events', headers=INGEST,
                              json={'event_type': 'like', 'timestamp': future})).status_code == 422
    assert (await client.post(f'{API}/lives/{live}/phones', headers=INGEST,
                              json={'phone': '123'})).status_code == 422
    comments = (await client.get(f'{MKT}/lives/{live}/comments', headers=READ)).json()['items']
    first = comments[0] if comments[0]['username'] == 'usuario123' else comments[1]
    # A timestamp without offset is Lima wall-clock time (UTC-5).
    assert first['created_at'].startswith('2026-10-02T00:23:10')
    assert first['created_at_lima'].startswith('2026-10-01T19:23:10')
    assert first['phone_detected'] == '+51987654321'


async def test_finish_consolidates_and_is_idempotent_and_blocks_late_events(env):
    client, _ = env
    live = await seed(client)
    await client.post(f'{API}/lives/{live}/comments', headers=INGEST,
                      json={'username': 'a', 'text': 'precio? 987654321'})
    done = await client.patch(f'{API}/lives/{live}/finish', headers=INGEST, json={})
    assert done.status_code == 200
    body = done.json()
    assert body['status'] == 'finished' and body['duration_seconds'] >= 7000
    assert body['metrics']['phones_captured'] == 1
    again = await client.patch(f'{API}/lives/{live}/finish', headers=INGEST)
    assert again.json()['ended_at'] == body['ended_at']
    stale = (datetime.now(timezone.utc) + timedelta(minutes=4)).isoformat()
    late = await client.post(f'{API}/lives/{live}/events', headers=INGEST,
                             json={'event_type': 'like', 'timestamp': stale})
    assert late.status_code == 200  # inside the 5-minute grace window
    filters = {'status': 'finished'}
    listed = (await client.get(f'{MKT}/lives', headers=READ, params=filters)).json()
    assert [i['id'] for i in listed['items']] == [live]


async def test_contact_time_and_contact_performance(env):
    client, factory = env
    live = await seed(client)
    base = datetime.now(timezone.utc) - timedelta(hours=1)
    ids = []
    for i, phone in enumerate(['987654321', '976543210', '965432109', '954321098']):
        res = await client.post(f'{API}/lives/{live}/phones', headers=INGEST, json={
            'phone': phone, 'tiktok_username': f'u{i}', 'source': 'manual_capture',
            'captured_at': (base + timedelta(minutes=i)).isoformat()})
        assert res.status_code == 201
        ids.append(res.json()['id'])
    dup = await client.post(f'{API}/lives/{live}/phones', headers=INGEST,
                            json={'phone': '+51 987 654 321'})
    assert dup.status_code == 200 and dup.json()['created'] is False

    # contacted after 2 min, 10 min, 40 min; the fourth stays new.
    for pid, minutes in zip(ids[:3], (2, 10, 40), strict=True):
        async with factory() as s, s.begin():
            captured = await s.scalar(select(m.CapturedPhone.captured_at).where(
                m.CapturedPhone.id == pid))
        patch = await client.patch(f'{API}/phones/{pid}', headers=INGEST, json={
            'status': 'contacted', 'advisor_id': 'asesor1',
            'contacted_at': (captured.replace(tzinfo=timezone.utc)
                             + timedelta(minutes=minutes)).isoformat()})
        assert patch.status_code == 200, patch.text
        assert patch.json()['contact_time_seconds'] == minutes * 60

    perf = (await client.get(f'{MKT}/contact-performance', headers=READ)).json()
    assert perf['phones'] == 4 and perf['contacted'] == 3 and perf['uncontacted'] == 1
    assert perf['min_seconds'] == 120 and perf['max_seconds'] == 2400
    assert perf['median_seconds'] == 600
    assert perf['average_seconds'] == pytest.approx((120 + 600 + 2400) / 3, abs=0.1)
    assert perf['percent'] == {'lt_5_min': 33.3, '5_to_15_min': 33.3,
                               '15_to_30_min': 0.0, 'gt_30_min': 33.3}
    listing = (await client.get(f'{MKT}/lives/{live}/phones', headers=READ,
                                params={'status': 'new'})).json()['items']
    assert len(listing) == 1 and listing[0]['contact_time_seconds'] is None
    early = await client.patch(f'{API}/phones/{ids[3]}', headers=INGEST, json={
        'status': 'contacted', 'contacted_at': (base - timedelta(days=1)).isoformat()})
    assert early.status_code == 422


async def test_hourly_ranking_accounts_and_overview(env):
    client, factory = env
    await seed(client, 'LIVE-A', 'medical', 'lucia')
    await client.post(f'{API}/accounts', headers=INGEST,
                      json={'id': 'medical-2', 'name': 'Medical 2', 'username': 'medical.cirugias2'})
    await client.post(f'{API}/tiktokers', headers=INGEST, json={'id': 'andrea', 'name': 'Andrea'})
    await client.post(f'{API}/lives', headers=INGEST, json={
        'live_id': 'LIVE-B', 'tiktoker_id': 'andrea', 'account_id': 'medical-2',
        'started_at': (datetime.now(timezone.utc) - timedelta(hours=4)).isoformat()})
    at = datetime(2026, 10, 14, 17, 30, tzinfo=timezone.utc)  # 12:30 Lima
    for live, phone in [('LIVE-A', '987654321'), ('LIVE-A', '976543210'),
                        ('LIVE-B', '965432109')]:
        stamp = at if live == 'LIVE-A' else at + timedelta(hours=2)
        r = await client.post(f'{API}/lives/{live}/phones', headers=INGEST, json={
            'phone': phone, 'captured_at': stamp.isoformat()})
        assert r.status_code == 201, r.text
    await client.post(f'{API}/lives/LIVE-A/events', headers=INGEST, json={
        'event_type': 'viewer', 'value': 80})

    hourly = (await client.get(f'{MKT}/hourly-performance', headers=READ)).json()
    by_hour = {i['hour']: i['phones'] for i in hourly['items']}
    assert by_hour[12] == 2 and by_hour[14] == 1 and hourly['best_hour'] == 12
    assert hourly['total'] == 3 and len(hourly['items']) == 24
    daily = (await client.get(f'{MKT}/hourly-performance', headers=READ,
                              params={'granularity': 'day'})).json()['items']
    assert daily == [{'period': '2026-10-14', 'phones': 3}]
    monthly = (await client.get(f'{MKT}/hourly-performance', headers=READ,
                                params={'granularity': 'month'})).json()['items']
    assert monthly == [{'period': '2026-10', 'phones': 3}]

    board = (await client.get(f'{MKT}/ranking', headers=READ,
                              params={'sort_by': 'phones'})).json()['items']
    assert [i['tiktoker_id'] for i in board] == ['lucia', 'andrea'] and board[0]['rank'] == 1
    assert (await client.get(f'{MKT}/ranking', headers=READ,
                             params={'sort_by': 'views'})).status_code == 200
    assert (await client.get(f'{MKT}/ranking', headers=READ,
                             params={'sort_by': 'bogus'})).status_code == 422

    accounts = (await client.get(f'{MKT}/accounts/performance', headers=READ)).json()['items']
    assert {a['account_id']: a['phones'] for a in accounts} == {'medical': 2, 'medical-2': 1}
    assert accounts[0]['account'] == '@medical.cirugias'

    overview = (await client.get(f'{MKT}/overview', headers=READ)).json()
    assert overview['lives'] == 2 and overview['phones'] == 3 and overview['live_now'] == 2
    assert overview['peak_viewers'] == 80 and overview['unique_viewers'] is None
    assert overview['leads_per_hour'] > 0
    only = (await client.get(f'{MKT}/overview', headers=READ,
                             params={'account_id': 'medical-2'})).json()
    assert only['lives'] == 1 and only['phones'] == 1


async def test_export_xlsx_in_memory_and_retention_keeps_history(env):
    client, factory = env
    live = await seed(client)
    for i, text in enumerate(['quiero información 987654321', '=HYPERLINK("http://x","clic")',
                              'hola buenas']):
        await client.post(f'{API}/lives/{live}/comments', headers=INGEST,
                          json={'username': f'user{i}', 'text': text, 'event_id': f'c{i}'})
    await client.post(f'{API}/lives/{live}/events', headers=INGEST,
                      json={'event_type': 'like', 'value': 9, 'event_id': 'like-1'})

    exported = await client.get(f'{MKT}/lives/{live}/comments/export', headers=EXPORT)
    assert exported.status_code == 200
    assert 'spreadsheetml' in exported.headers['content-type']
    assert 'comentarios_LIVE-1025.xlsx' in exported.headers['content-disposition']
    sheet = load_workbook(io.BytesIO(exported.content)).active
    rows = list(sheet.values)
    assert rows[0] == ('fecha', 'hora', 'usuario TikTok', 'comentario', 'teléfono detectado',
                       'live', 'TikToker', 'cuenta')
    assert len(rows) == 4
    assert rows[1][2:] == ('user0', 'quiero información 987654321', '+51987654321',
                           'LIVE-1025', 'Lucía Ramos', '@medical.cirugias')
    formula_cell = list(sheet.iter_rows(min_row=3, max_row=3))[0][3]
    assert formula_cell.value.startswith('=HYPERLINK') and formula_cell.data_type == 's'

    # Age the temporary rows beyond the retention window and run the daily job.
    old = datetime.now(timezone.utc) - timedelta(days=45)
    async with factory() as s, s.begin():
        await s.execute(update(m.LiveComment).values(created_at=old))
        await s.execute(update(m.LiveEventKey).values(created_at=old))
        await s.execute(update(m.Live).values(
            started_at=old - timedelta(hours=2), ended_at=old, status='finished',
            duration_seconds=7200))
    removed = await purge_expired(factory)
    assert removed['comments'] == 3 and removed['event_keys'] == 4

    gone = await client.get(f'{MKT}/lives/{live}/comments/export', headers=EXPORT)
    assert gone.status_code == 410
    assert gone.json()['detail']['code'] == 'COMMENTS_RETENTION_EXPIRED'
    assert (await client.get(f'{MKT}/lives/{live}/comments', headers=READ)).status_code == 410

    # Historical data survives: lives, metrics and phones.
    detail = (await client.get(f'{MKT}/lives/{live}', headers=READ)).json()
    assert detail['comments_available'] is False
    assert detail['metrics']['comments_count'] == 3 and detail['metrics']['likes_count'] == 9
    assert detail['phones']['total'] == 1
    async with factory() as s:
        assert await s.scalar(select(func.count()).select_from(m.Live)) == 1
        assert await s.scalar(select(func.count()).select_from(m.CapturedPhone)) == 1
        assert await s.scalar(select(func.count()).select_from(m.LiveMetrics)) == 1
