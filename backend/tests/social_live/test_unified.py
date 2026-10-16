"""Behaviour of the unified flow: one contact per number, first-capture reporting, retention
policy, attention workflow, data conversion and the public routes."""
import asyncio
import re
from datetime import datetime, timedelta, timezone
from pathlib import Path

import pytest
from fastapi import HTTPException
from sqlalchemy import create_engine, func, select, text, update
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from livepulse.modules.analytics.models import (
    EventFact,
    Lead,
    LeadAlias,
    LeadAudit,
    LeadPhone,
    UserRole,
)
from livepulse.modules.broadcasts.models import Broadcast, RawEvent
from livepulse.modules.social_live import backfill, contacts
from livepulse.modules.social_live import models as m
from livepulse.shared import auth as clerk
from livepulse.shared.config import settings
from livepulse.shared.db import Base

INGEST = {'Authorization': 'Bearer ingest-secret'}
READ = {'Authorization': 'Bearer read-secret'}
EXPORT = {'Authorization': 'Bearer export-secret'}
API, MKT = '/api/social-live', '/api/marketing/tiktok'


# What Clerk's Public metadata would say for each test person (boss = ADMIN_USER_IDS).
CLERK_ROLES = {'ana': 'agente_ventas', 'bea': 'agente_ventas', 'sup': 'supervisor',
               'mk': 'jefe_marketing', 'jefa': 'jefe_marketing'}


def person(name):
    return {'Authorization': f'Bearer user:{name}'}


@pytest.fixture
def people(monkeypatch):
    monkeypatch.setattr(settings, 'admin_user_ids', ['boss'])
    monkeypatch.setattr(clerk, 'clerk_configured', lambda: True)

    def fake_verify(token):
        if not token.startswith('user:'):
            raise HTTPException(401, 'Sesión inválida o vencida.')
        uid = token.removeprefix('user:')
        return {'id': uid, 'role': CLERK_ROLES.get(uid)}

    monkeypatch.setattr(clerk, 'verify_token', fake_verify)


async def account(client, account_id, username):
    r = await client.post(f'{API}/accounts', headers=INGEST,
                          json={'id': account_id, 'name': account_id, 'username': username})
    assert r.status_code == 200, r.text


async def live(client, live_id, account_id='medical', started=None):
    r = await client.post(f'{API}/lives', headers=INGEST, json={
        'live_id': live_id, 'account_id': account_id,
        'started_at': (started or datetime.now(timezone.utc) - timedelta(hours=3)).isoformat()})
    assert r.status_code == 201, r.text


async def phone(client, live_id, number, at, user='u'):
    r = await client.post(f'{API}/lives/{live_id}/phones', headers=INGEST, json={
        'phone': number, 'tiktok_username': user, 'captured_at': at.isoformat()})
    assert r.status_code in (200, 201), r.text
    return r.json()


# ---------- contacts: one number, one contact ----------

async def test_same_number_in_two_accounts_and_lives_is_one_new_contact(env):
    client, factory = env
    await account(client, 'medical', 'medical.cirugias')
    await account(client, 'medical-2', 'medical.cirugias2')
    t0 = datetime.now(timezone.utc) - timedelta(hours=5)
    await live(client, 'A1', 'medical', t0)
    await live(client, 'B1', 'medical-2', t0 + timedelta(hours=1))
    await live(client, 'A2', 'medical', t0 + timedelta(hours=2))
    first = await phone(client, 'A1', '987654321', t0 + timedelta(minutes=10), 'maria')
    second = await phone(client, 'B1', '+51 987 654 321', t0 + timedelta(hours=1, minutes=5))
    third = await phone(client, 'A2', '987-654-321', t0 + timedelta(hours=2, minutes=5))
    assert (first['outcome'], second['outcome'], third['outcome']) == ('new', 'mention', 'mention')
    assert first['id'] == second['id'] == third['id']  # one contact, three appearances

    overview = (await client.get(f'{MKT}/overview', headers=READ)).json()
    assert overview['phones'] == 1 and overview['phone_mentions'] == 3
    accounts = {a['account_id']: a['phones']
                for a in (await client.get(f'{MKT}/accounts/performance',
                                           headers=READ)).json()['items']}
    assert accounts.get('medical') == 1 and accounts.get('medical-2', 0) == 0  # first capture
    async with factory() as s:
        assert await s.scalar(select(func.count()).select_from(m.PhoneRegistry)) == 1
        assert await s.scalar(select(func.count()).select_from(m.CapturedPhone)) == 3


async def test_a_late_event_with_an_earlier_time_becomes_the_first_capture(env):
    client, _ = env
    await account(client, 'medical', 'medical.cirugias')
    await account(client, 'medical-2', 'medical.cirugias2')
    t0 = datetime.now(timezone.utc) - timedelta(hours=4)
    await live(client, 'A', 'medical', t0)
    await live(client, 'B', 'medical-2', t0)
    later = await phone(client, 'B', '987654321', t0 + timedelta(hours=2))
    earlier = await phone(client, 'A', '987654321', t0 + timedelta(minutes=30))
    assert later['outcome'] == 'new' and earlier['outcome'] == 'new'  # earlier one wins
    perf = {a['account_id']: a['phones']
            for a in (await client.get(f'{MKT}/accounts/performance', headers=READ)).json()['items']}
    assert perf.get('medical') == 1 and perf.get('medical-2', 0) == 0
    a = (await client.get(f'{MKT}/lives/A/metrics', headers=READ)).json()
    b = (await client.get(f'{MKT}/lives/B/metrics', headers=READ)).json()
    assert (a['phones_captured'], a['phone_mentions']) == (1, 1)
    assert (b['phones_captured'], b['phone_mentions']) == (0, 1)


async def test_simultaneous_first_captures_create_one_contact(tmp_path, monkeypatch):
    """Two writers race on the same number: the unique contact decides, nobody errors."""
    from livepulse.modules.social_live import service

    engine = create_async_engine(f'sqlite+aiosqlite:///{tmp_path / "race.db"}')
    async with engine.begin() as connection:
        await connection.run_sync(lambda sync: Base.metadata.create_all(sync, tables=[
            m.TiktokAccount.__table__, m.Tiktoker.__table__, m.Live.__table__,
            m.LiveMetrics.__table__, m.PhoneRegistry.__table__, m.CapturedPhone.__table__,
            m.ContactAudit.__table__]))
    factory = async_sessionmaker(engine, expire_on_commit=False)
    now = datetime.now(timezone.utc)
    async with factory() as s, s.begin():
        s.add(m.TiktokAccount(id='medical', name='m', username='m'))
        for lid in ('L1', 'L2'):
            s.add(m.Live(id=lid, account_id='medical', started_at=now - timedelta(hours=1),
                         status='live'))
            s.add(m.LiveMetrics(live_id=lid))

    async def capture(lid):
        async with factory() as s, s.begin():
            lv = await s.get(m.Live, lid)
            return (await service.save_phone(s, lv, '+51987654321', 'u', 'comment', now))[0]

    results = await asyncio.gather(*(capture(lid) for lid in ('L1', 'L2', 'L1', 'L2')))
    assert sorted(r for r in results if r) == ['mention', 'mention', 'new'] or \
        results.count('new') == 1
    async with factory() as s:
        assert await s.scalar(select(func.count()).select_from(m.PhoneRegistry)) == 1
        firsts = await s.scalar(select(func.count()).select_from(m.CapturedPhone).where(
            m.CapturedPhone.is_new.is_(True)))
        assert firsts == 1
    await engine.dispose()


async def test_first_capture_date_uses_the_lima_calendar_across_midnight(env):
    client, _ = env
    await account(client, 'medical', 'medical.cirugias')
    # 23:50 in Lima on the 10th is 04:50 UTC on the 11th.
    lima_night = datetime(2026, 10, 11, 4, 50, tzinfo=timezone.utc)
    await live(client, 'N', started=lima_night - timedelta(hours=1))
    await phone(client, 'N', '987654321', lima_night)
    await phone(client, 'N', '976543210', lima_night + timedelta(minutes=20))  # 00:10 on the 11th

    day10 = (await client.get(f'{MKT}/overview', headers=READ, params={
        'date_from': '2026-10-10', 'date_to': '2026-10-10'})).json()
    day11 = (await client.get(f'{MKT}/overview', headers=READ, params={
        'date_from': '2026-10-11', 'date_to': '2026-10-11'})).json()
    assert day10['phones'] == 1 and day11['phones'] == 1
    daily = (await client.get(f'{MKT}/hourly-performance', headers=READ,
                              params={'granularity': 'day'})).json()['items']
    assert daily == [{'period': '2026-10-10', 'phones': 1}, {'period': '2026-10-11', 'phones': 1}]
    hours = {i['hour']: i['phones'] for i in (await client.get(
        f'{MKT}/hourly-performance', headers=READ)).json()['items']}
    assert hours[23] == 1 and hours[0] == 1
    month = (await client.get(f'{MKT}/hourly-performance', headers=READ,
                              params={'granularity': 'month'})).json()['items']
    assert month == [{'period': '2026-10', 'phones': 2}]


# ---------- retention policy ----------

async def test_resending_after_cleanup_never_changes_metrics(env):
    client, factory = env
    await account(client, 'medical', 'medical.cirugias')
    await live(client, 'OPEN')
    event = {'event_type': 'like', 'value': 7, 'event_id': 'lk-1'}
    await client.post(f'{API}/lives/OPEN/events', headers=INGEST, json=event)

    from livepulse.modules.social_live.retention import purge_expired
    old = datetime.now(timezone.utc) - timedelta(days=45)
    async with factory() as s, s.begin():
        await s.execute(update(m.LiveEventKey).values(created_at=old))
        await s.execute(update(m.LiveComment).values(created_at=old))
    await purge_expired(factory)
    again = await client.post(f'{API}/lives/OPEN/events', headers=INGEST, json=event)
    assert again.json()['duplicates'] == 1
    # Explicit policy for anything older than the window: rejected, never summed.
    stale = await client.post(f'{API}/lives/OPEN/events', headers=INGEST, json={
        'event_type': 'like', 'value': 5, 'event_id': 'lk-old', 'timestamp': old.isoformat()})
    assert stale.status_code == 409 and stale.json()['detail']['code'] == 'EVENT_TOO_OLD'
    metrics = (await client.get(f'{MKT}/lives/OPEN/metrics', headers=READ)).json()
    assert metrics['likes_count'] == 7


async def test_monitor_history_is_converted_once_even_after_cleanup(env):
    client, factory = env
    from livepulse.modules.social_live.publisher import publish_batch
    from livepulse.modules.social_live.retention import purge_expired
    now = datetime.now(timezone.utc).replace(microsecond=0)
    async with factory() as s, s.begin():
        s.add(Broadcast(id=1, account_id='medical', tiktok_username='medical.cirugias',
                        room_id='r', started_at=now - timedelta(hours=2), end_confirmed=False))
        s.add(RawEvent(event_id='medical:r:like:1', account_id='medical', broadcast_id=1,
                       type='like', occurred_at=now - timedelta(minutes=5),
                       payload={'count': 9}))
        s.add(RawEvent(event_id='medical:r:comment:1', account_id='medical', broadcast_id=1,
                       type='comment', occurred_at=now - timedelta(minutes=4),
                       payload={'user': '5', 'username': 'ana'}, raw_text='hola 987654321'))
    assert await publish_batch(factory) == 2
    old = now - timedelta(days=60)
    async with factory() as s, s.begin():
        await s.execute(update(m.LiveEventKey).values(created_at=old))
        await s.execute(update(m.Live).values(status='finished', started_at=old - timedelta(hours=1),
                                              ended_at=old, duration_seconds=3600))
    await purge_expired(factory)        # keys of the finished live are gone...
    assert await publish_batch(factory) == 0    # ...and nothing is re-counted
    metrics = (await client.get(f'{MKT}/lives/TT-1/metrics', headers=READ)).json()
    assert metrics['likes_count'] == 9 and metrics['comments_count'] == 1
    assert metrics['phones_captured'] == 1


# ---------- capture -> processing -> API -> visible ----------

async def test_captured_event_is_visible_to_medical360_and_the_old_bridge(env):
    client, factory = env
    from livepulse.modules.social_live.publisher import publish_batch
    now = datetime.now(timezone.utc).replace(microsecond=0)
    async with factory() as s, s.begin():
        s.add(Broadcast(id=3, account_id='medical', tiktok_username='medical.cirugias',
                        room_id='r3', started_at=now - timedelta(hours=1), end_confirmed=False))
        s.add(RawEvent(event_id='medical:r3:comment:1', account_id='medical', broadcast_id=3,
                       type='comment', occurred_at=now - timedelta(minutes=3),
                       payload={'user': '9', 'username': 'maria'},
                       raw_text='me interesa, mi número es 985//857#968'))
        s.add(RawEvent(event_id='medical:r3:audience:1', account_id='medical', broadcast_id=3,
                       type='audience', occurred_at=now - timedelta(minutes=2),
                       payload={'viewers': 30}))
    await publish_batch(factory)

    monitor = (await client.get(f'{MKT}/monitor', headers=READ)).json()
    row = next(a for a in monitor['accounts'] if a['account_id'] == 'medical')
    assert row['on_air'] and row['tiktoker'] is None
    assert row['metrics']['peak_viewers'] == 30 and row['metrics']['unique_viewers'] is None
    assert row['metrics']['direct_messages_capture'] is False
    inbox = (await client.get(f'{MKT}/inbox', headers=READ)).json()['items']
    assert inbox[0]['phones'] == ['+51985857968'] and inbox[0]['tiktoker'] == 'Sin asignar'
    assert inbox[0]['state'] == 'Por contactar' and inbox[0]['is_first_capture'] is True
    numbers = (await client.get(f'{MKT}/numbers', headers=READ)).json()
    assert numbers['received'] == 1 and numbers['phones'][0]['phone'] == '+51985857968'
    assert numbers['from_dm'] is None  # no private-message source connected
    # A marketing person never gets the contact list.
    assert (await client.get(f'{MKT}/numbers', headers={'Authorization': 'Bearer nope'})
            ).status_code == 401

    legacy = (await client.get('/v1/integrations/medical360/inbox', headers={
        'X-Medical360-Token': 'read-secret'}))
    assert legacy.status_code == 200 and legacy.json()['items'][0]['phones'] == ['+51985857968']


# ---------- attention workflow ----------

async def test_contact_workflow_is_persisted_audited_and_race_safe(env, people):
    client, factory = env
    await account(client, 'medical', 'medical.cirugias')
    await live(client, 'L')
    async with factory() as s, s.begin():
        s.add_all([UserRole(user_id='ana', role='agente_ventas', display_name='Ana'),
                   UserRole(user_id='bea', role='agente_ventas', display_name='Bea'),
                   UserRole(user_id='sup', role='supervisor', display_name='Sara'),
                   UserRole(user_id='mk', role='jefe_marketing', display_name='Jefa')])
    contact = await phone(client, 'L', '987654321', datetime.now(timezone.utc) - timedelta(hours=1))
    cid = contact['id']

    # Authorization: marketing and service tokens cannot work contacts; agents can.
    assert (await client.patch(f'/api/leads/{cid}', headers=person('mk'),
                               json={'action': 'claim'})).status_code == 403
    assert (await client.patch(f'/api/leads/{cid}', headers=INGEST,
                               json={'action': 'claim'})).status_code == 403
    assert (await client.patch(f'/api/leads/{cid}', headers=person('ana'),
                               json={'action': 'claim'})).status_code == 200
    # Bea saw it free; Ana won: a clear conflict, and Bea no longer sees it.
    lost = await client.patch(f'/api/leads/{cid}', headers=person('bea'),
                              json={'action': 'claim', 'expect_unassigned': True})
    assert lost.status_code == 409
    assert (await client.get('/api/inbox', headers=person('bea'),
                             params={'phone_only': True})).json()['items'] == []
    sup_view = (await client.get('/api/inbox', headers=person('sup'),
                                 params={'phone_only': True})).json()['items']
    assert sup_view[0]['owner'] == 'Ana'

    # The row-level guard: a stale write (someone changed it after we read it) is refused.
    async with factory() as s:
        stale = await s.get(m.PhoneRegistry, cid)
        stale_status, stale_advisor = stale.status, stale.advisor_id
        async with factory() as other, other.begin():
            await other.execute(update(m.PhoneRegistry).where(m.PhoneRegistry.id == cid)
                                .values(advisor_id='bea'))
        stale.advisor_id, stale.status = stale_advisor, stale_status
        with pytest.raises(HTTPException) as caught:
            await contacts._apply(s, stale, {'status': 'contacted'})
        assert caught.value.status_code == 409
    async with factory() as s, s.begin():
        await s.execute(update(m.PhoneRegistry).where(m.PhoneRegistry.id == cid)
                        .values(advisor_id='ana'))

    # First contact time is recorded once and survives re-opening.
    done = await client.patch(f'/api/leads/{cid}', headers=person('ana'),
                              json={'status': 'contacted'})
    first_contact = done.json()['contactedAt']
    assert first_contact and done.json()['contactTimeSeconds'] >= 3500
    await client.patch(f'/api/leads/{cid}', headers=person('ana'), json={'status': 'new'})
    again = await client.patch(f'/api/leads/{cid}', headers=person('ana'),
                               json={'status': 'qualified'})
    assert again.json()['contactedAt'] == first_contact

    # Visible to others after a reload, and every change names who made it.
    seen = (await client.get('/api/inbox', headers=person('sup'),
                             params={'phone_only': True})).json()['items'][0]
    assert seen['statusCode'] == 'qualified' and seen['owner'] == 'Ana'
    log = (await client.get(f'/api/leads/{cid}/history', headers=person('sup'))).json()['items']
    assert [(i['actor'], i['action']) for i in log if i['action'] != 'captured'] == [
        ('ana', 'claim'), ('ana', 'status'), ('ana', 'status'), ('ana', 'status')]
    assert (await client.get(f'/api/leads/{cid}/history', headers=person('ana'))).status_code == 403


# ---------- export of the selected live ----------

async def test_export_contains_only_the_requested_live_and_reports_expired_retention(env):
    client, factory = env
    from openpyxl import load_workbook
    await account(client, 'medical', 'medical.cirugias')
    await live(client, 'LIVE-A')
    await live(client, 'LIVE-B')
    for live_id, text_ in (('LIVE-A', 'comentario de A'), ('LIVE-B', 'comentario de B')):
        await client.post(f'{API}/lives/{live_id}/comments', headers=INGEST,
                          json={'username': 'u', 'text': text_})
    sheet = load_workbook(__import__('io').BytesIO((await client.get(
        f'{MKT}/lives/LIVE-A/comments/export', headers=EXPORT)).content)).active
    values = [row[3] for row in sheet.iter_rows(min_row=2, values_only=True)]
    assert values == ['comentario de A']
    old = datetime.now(timezone.utc) - timedelta(days=45)
    async with factory() as s, s.begin():
        await s.execute(update(m.LiveComment).where(m.LiveComment.live_id == 'LIVE-B')
                        .values(created_at=old))
        await s.execute(update(m.Live).where(m.Live.id == 'LIVE-B').values(
            status='finished', started_at=old - timedelta(hours=1), ended_at=old))
    from livepulse.modules.social_live.retention import purge_expired
    await purge_expired(factory)
    expired = await client.get(f'{MKT}/lives/LIVE-B/comments/export', headers=EXPORT)
    assert expired.status_code == 410
    assert expired.json()['detail']['code'] == 'COMMENTS_RETENTION_EXPIRED'
    assert (await client.get(f'{MKT}/lives/LIVE-A/comments/export',
                             headers=EXPORT)).status_code == 200


# ---------- honest metrics ----------

async def test_unavailable_metrics_are_explicit(env):
    client, _ = env
    await account(client, 'medical', 'medical.cirugias')
    await live(client, 'L')
    await client.post(f'{API}/lives/L/events', headers=INGEST,
                      json={'event_type': 'direct_message', 'value': 3})
    metrics = (await client.get(f'{MKT}/lives/L/metrics', headers=READ)).json()
    assert metrics['unique_viewers'] is None and metrics['average_viewers'] is None
    assert metrics['direct_messages_reported'] == 3
    assert metrics['direct_messages_capture'] is False
    await client.post(f'{API}/lives/L/events', headers=INGEST,
                      json={'event_type': 'unique_viewers', 'value': 120})
    metrics = (await client.get(f'{MKT}/lives/L/metrics', headers=READ)).json()
    assert metrics['unique_viewers'] == 120 and metrics['peak_viewers'] == 0


# ---------- conversion of previous data (migration 0006) ----------

def legacy_database():
    """Release-0005 shape: appearances carry their own status/advisor/contact time."""
    engine = create_engine('sqlite://')
    Base.metadata.create_all(engine, tables=[
        m.TiktokAccount.__table__, m.Tiktoker.__table__, m.Live.__table__,
        m.LiveMetrics.__table__, m.PhoneRegistry.__table__, Broadcast.__table__,
        RawEvent.__table__, m.LiveEventKey.__table__, m.LiveInterestedUser.__table__, EventFact.__table__, Lead.__table__,
        LeadPhone.__table__, LeadAlias.__table__, LeadAudit.__table__, UserRole.__table__])
    with engine.begin() as conn:
        conn.execute(text(
            'CREATE TABLE captured_phones (id INTEGER PRIMARY KEY, live_id VARCHAR(40), '
            'phone VARCHAR(20), tiktok_username VARCHAR(128), source VARCHAR(24), '
            'captured_at TIMESTAMP, contacted_at TIMESTAMP, advisor_id VARCHAR(64), '
            "status VARCHAR(16) DEFAULT 'new', is_new BOOLEAN DEFAULT 1)"))
    return engine


def test_previous_data_is_converted_without_loss_or_duplicates():
    engine = legacy_database()
    t = datetime(2026, 10, 1, 15, 0)
    with engine.begin() as conn:
        conn.execute(text(
            'INSERT INTO broadcasts (id, account_id, tiktok_username, room_id, started_at, '
            'ended_at, end_confirmed) VALUES '
            "(1, 'medical', 'medical.cirugias', 'r1', :t, NULL, 0), "
            "(2, 'medical-2', 'medical.cirugias2', 'r2', :t2, :t3, 1)"),
            {'t': t, 't2': t + timedelta(hours=3), 't3': t + timedelta(hours=4)})
        # Release 0005 published some lives itself.
        conn.execute(text("INSERT INTO tiktok_accounts (id, name, username) "
                          "VALUES ('medical', 'medical', 'medical.cirugias')"))
        conn.execute(text(
            "INSERT INTO lives (id, account_id, started_at, status) VALUES "
            "('TT-1', 'medical', :t, 'live')"), {'t': t})
        conn.execute(text("INSERT INTO live_metrics (live_id) VALUES ('TT-1')"))
        for row in [(1, 'TT-1', '+51987654321', 'maria', 'comment', t + timedelta(minutes=10),
                     None, None, 'new'),
                    (2, 'TT-1', '+51976543210', 'luis', 'comment', t + timedelta(minutes=20),
                     None, None, 'new'),
                    (3, 'TT-2', '+51987654321', 'maria', 'comment', t + timedelta(hours=3, minutes=5),
                     t + timedelta(hours=3, minutes=30), 'ana', 'contacted')]:
            conn.execute(text(
                'INSERT INTO captured_phones (id, live_id, phone, tiktok_username, source, '
                'captured_at, contacted_at, advisor_id, status) VALUES '
                '(:a, :b, :c, :d, :e, :f, :g, :h, :i)'),
                dict(zip('abcdefghi', row, strict=True)))
        # previous analytics model: a contacted lead with history, a repeated and a manual one
        conn.execute(text("INSERT INTO leads (id, status, review_needed, created_at) VALUES "
                          "('L1', 'contacted', 0, :t), ('L2', 'new', 0, :t), ('L3', 'ready', 0, :t)"),
                     {'t': t})
        conn.execute(text("INSERT INTO lead_phones (phone, lead_id) VALUES "
                          "('+51965432109', 'L1'), ('+51987654321', 'L2'), ('+51954321098', 'L3')"))
        conn.execute(text("INSERT INTO lead_aliases (user_key, lead_id, username) "
                          "VALUES ('77', 'L1', 'carla')"))
        conn.execute(text(
            'INSERT INTO raw_events (event_id, account_id, broadcast_id, type, occurred_at, '
            "payload) VALUES ('e-pub', 'medical', 1, 'like', :t, '{}'), "
            "('e-new', 'medical', 1, 'like', :t, '{}')"), {'t': t})
        conn.execute(text("INSERT INTO event_facts (event_id, broadcast_id, account_id, "
                          "user_key, kind, occurred_at, comments, likes, shares, gifts, diamonds, "
                          "interested, phones, lead_id) VALUES ('e-pub', 1, 'medical', '77', "
                          "'comment', :t, 1, 0, 0, 0, 0, 0, '[]', 'L1')"), {'t': t})
        conn.execute(text("INSERT INTO live_event_keys (live_id, event_id) "
                          "VALUES ('TT-1', 'e-pub')"))
        conn.execute(text(
            "INSERT INTO lead_audit (id, lead_id, actor, action, previous, value, created_at) "
            "VALUES ('a1', 'L1', 'x', 'status', 'new', 'contacted', :t)"),
            {'t': t + timedelta(minutes=45)})

    with engine.begin() as conn:
        first = backfill.convert_all(conn)
    with engine.begin() as conn:
        second = backfill.convert_all(conn)

    assert first['lives'] == 1 and first['contacts'] == 2 and first['legacy_contacts'] == 2
    assert second == {'lives': 0, 'contacts': 0, 'legacy_contacts': 0, 'published': 0}
    with engine.connect() as conn:
        registry = {r.phone: r for r in conn.execute(text('SELECT * FROM phone_registry'))}
        assert set(registry) == {'+51987654321', '+51976543210', '+51965432109', '+51954321098'}
        repeated = registry['+51987654321']
        # L2 was registered manually before either live appearance; keep that history.
        assert repeated.first_live_id is None and repeated.first_account_id is None
        assert backfill._moment(repeated.first_captured_at) == t
        assert repeated.status == 'contacted' and repeated.advisor_id == 'ana'
        assert repeated.contacted_at is not None
        carla = registry['+51965432109']
        assert carla.status == 'contacted' and carla.first_username == 'carla'
        assert carla.first_live_id == 'TT-1' and carla.contacted_at is not None
        manual = registry['+51954321098']
        assert manual.status == 'new' and manual.first_live_id is None
        assert manual.first_source == 'manual_capture'
        news = [r.id for r in conn.execute(text('SELECT id FROM captured_phones WHERE is_new = 1'))]
        assert sorted(news) == [2]
        metrics = conn.execute(text(
            "SELECT phones_captured, phone_mentions FROM live_metrics WHERE live_id = 'TT-1'")).one()
        assert tuple(metrics) == (1, 2)
        published = {r[0]: r[1] for r in conn.execute(text(
            'SELECT event_id, published_at FROM raw_events'))}
        assert published['e-pub'] is not None and published['e-new'] is None
        assert conn.execute(text('SELECT COUNT(*) FROM lives')).scalar() == 2  # TT-1 + TT-2


# ---------- public routes through the proxy ----------

def nginx_locations(path: Path):
    """Parse `location` blocks: [(kind, prefix, proxy_pass)]."""
    text_ = path.read_text(encoding='utf-8')
    found = []
    for kind, prefix, body in re.findall(r'location\s+(=\s+)?(\S+)\s*\{([^}]*)\}', text_):
        match = re.search(r'proxy_pass\s+http://[^/\s;]+(/[^;\s]*)?;', body)
        if match:
            found.append((bool(kind), prefix, match.group(1) or '/'))
    return found


def through_proxy(locations, public_path):
    """nginx rule: exact match wins, else the longest prefix; the matched prefix is replaced
    by the URI given in proxy_pass."""
    for exact, prefix, target in locations:
        if exact and public_path == prefix:
            return target
    candidates = [(len(p), p, t) for exact, p, t in locations
                  if not exact and public_path.startswith(p)]
    if not candidates:
        return None
    _, prefix, target = max(candidates)
    return target + public_path[len(prefix):]


@pytest.mark.parametrize('conf', ['apps/social-pulse/nginx.conf'])
async def test_public_api_paths_reach_existing_backend_routes(env, conf):
    client, _ = env
    locations = nginx_locations(Path(__file__).resolve().parents[3] / conf)
    assert locations, 'no proxy locations parsed'
    expected_public = {
        '/api/health': 200,
        '/api/v1/shifts/active': 404,           # attendance has been retired
        '/api/auth/me': 401,                    # new routes keep their /api prefix
        '/api/inbox': 401,
        '/api/leads/summary': 401,
        '/api/team': 401,
        '/api/marketing/tiktok/overview': 401,
        '/api/social-live/lives': 401,
    }
    for public, want in expected_public.items():
        upstream = through_proxy(locations, public)
        assert upstream is not None, public
        method = 'POST' if upstream.endswith('/lives') else 'GET'
        response = await client.request(method, upstream)
        assert response.status_code == want, (conf, public, upstream, response.status_code)
    # The prefix must not be duplicated or dropped by the proxy.
    assert through_proxy(locations, '/api/auth/me') == '/api/auth/me'
    assert through_proxy(locations, '/api/v1/shifts/active') == '/v1/shifts/active'
    assert through_proxy(locations, '/api/health') == '/health'


# ---------- regression cases from the integration review ----------

@pytest.mark.parametrize('endpoint,field', [
    ('overview', 'phones'), ('numbers', 'received'),
    ('hourly-performance', 'total'), ('contact-performance', 'phones'),
])
async def test_date_bounds_are_independent(env, endpoint, field):
    client, _ = env
    await account(client, 'medical', 'medical.cirugias')
    start = (datetime.now(timezone.utc) - timedelta(days=5)).replace(
        hour=5, minute=0, second=0, microsecond=0)
    await live(client, 'DATES', started=start)
    for day, number in enumerate(('987654321', '976543210', '965432109')):
        await phone(client, 'DATES', number, start + timedelta(days=day, hours=1))
    url = f'{MKT}/{endpoint}'
    for params, expected in [
        ({'date_from': start.isoformat()}, 3),
        ({'date_to': (start + timedelta(days=1)).isoformat()}, 1),
        ({'date_from': start.isoformat(),
          'date_to': (start + timedelta(days=2)).isoformat()}, 2),
        ({'date_from': start.date().isoformat(),
          'date_to': (start + timedelta(days=2)).date().isoformat()}, 3),
    ]:
        result = await client.get(url, headers=READ, params=params)
        assert result.status_code == 200, result.text
        assert result.json()[field] == expected, (params, result.json())
    invalid = await client.get(url, headers=READ, params={
        'date_from': (start + timedelta(days=1)).isoformat(), 'date_to': start.isoformat()})
    assert invalid.status_code == 422


async def test_phone_filter_and_server_pagination_do_not_hide_pending_contacts(env, people):
    client, factory = env
    await account(client, 'medical', 'medical.cirugias')
    await live(client, 'INBOX')
    now = datetime.now(timezone.utc)
    first = await phone(client, 'INBOX', '987654321', now - timedelta(hours=1))
    async with factory() as s, s.begin():
        s.add(UserRole(user_id='ana', role='agente_ventas'))
        for i in range(201):
            s.add(m.LiveComment(live_id='INBOX', username=f'u{i}', comment='hola',
                               created_at=now - timedelta(seconds=i)))
        # A hidden/discarded contact must not consume a visible pagination slot.
        s.add(m.PhoneRegistry(phone='+51976543210', first_live_id='INBOX',
                             first_account_id='medical', first_captured_at=now,
                             first_source='comment', status='discarded'))
    headers = person('ana')
    params = {'limit': 200, 'exclude_discarded': True}
    phones = (await client.get('/api/inbox', headers=headers,
                              params={**params, 'phone_only': True})).json()
    assert [i['id'] for i in phones['items']] == [first['id']]
    assert phones['total'] == 1 and phones['next_offset'] is None
    page1 = (await client.get('/api/inbox', headers=headers, params=params)).json()
    page2 = (await client.get('/api/inbox', headers=headers,
                             params={**params, 'offset': page1['next_offset']})).json()
    assert page1['total'] == 202 and len(page1['items']) == 200
    assert len(page2['items']) == 2 and page2['next_offset'] is None
    ids = [i['id'] for i in page1['items'] + page2['items']]
    assert len(set(ids)) == 202 and first['id'] in ids


async def test_historical_interested_viewers_survive_retention_and_worker_restarts(env):
    from livepulse.modules.social_live.publisher import publish_batch
    from livepulse.modules.social_live.retention import purge_expired
    client, factory = env
    old = datetime.now(timezone.utc) - timedelta(days=45)
    async with factory() as s, s.begin():
        s.add(Broadcast(id=909, account_id='medical', tiktok_username='medical.cirugias',
                        room_id='history', started_at=old, ended_at=old + timedelta(hours=1),
                        end_confirmed=True))
        await s.flush()
        for i, user in enumerate(('ana', 'ana', 'bea', 'ana')):
            s.add(RawEvent(event_id=f'history-{i}', account_id='medical', broadcast_id=909,
                           type='comment', occurred_at=old + timedelta(minutes=i),
                           payload={'username': user}, raw_text='quiero informacion del precio'))
    assert await publish_batch(factory, limit=1) == 1
    await purge_expired(factory)  # temporary event keys can be deleted for this old live
    assert await publish_batch(factory) == 3
    assert await publish_batch(factory) == 0
    metrics = (await client.get(f'{MKT}/lives/TT-909/metrics', headers=READ)).json()
    assert metrics['comments_count'] == 4 and metrics['interested_viewers'] == 2
    async with factory() as s:
        assert await s.scalar(select(func.count()).select_from(m.LiveComment)) == 0
        assert await s.scalar(select(func.count()).select_from(m.LiveInterestedUser)) == 2


def test_migration_merges_legacy_attention_and_earliest_capture():
    engine = legacy_database()
    early = datetime(2026, 10, 1, 15)
    later = early + timedelta(days=1)
    with engine.begin() as conn:
        conn.execute(text("INSERT INTO tiktok_accounts(id,name,username) "
                          "VALUES ('medical','Medical','medical.cirugias')"))
        conn.execute(text("INSERT INTO lives(id,account_id,started_at,status) "
                          "VALUES ('L','medical',:t,'live')"), {'t': later})
        conn.execute(text("INSERT INTO live_metrics(live_id) VALUES ('L')"))
        conn.execute(text("INSERT INTO captured_phones(id,live_id,phone,source,captured_at,"
                          "status,advisor_id) VALUES "
                          "(1,'L','+51987654321','comment',:t,'new','current-advisor')"), {'t': later})
        conn.execute(text("INSERT INTO leads(id,status,review_needed,created_at) "
                          "VALUES ('OLD','contacted',0,:t)"), {'t': early})
        conn.execute(text("INSERT INTO lead_phones(phone,lead_id) "
                          "VALUES ('+51987654321','OLD')"))
        contacted = early + timedelta(minutes=15)
        conn.execute(text("INSERT INTO lead_audit(id,lead_id,actor,action,value,created_at) "
                          "VALUES ('a','OLD','legacy-agent','status','contacted',:t)"),
                     {'t': contacted})
        backfill.convert_all(conn)
        backfill.convert_all(conn)  # conversion is idempotent
        row = conn.execute(text('SELECT * FROM phone_registry')).one()
        assert row.status == 'contacted'
        assert backfill._moment(row.first_captured_at) == early
        assert backfill._moment(row.contacted_at) == contacted
        assert row.advisor_id == 'current-advisor'
        assert row.first_live_id is None  # no invented LIVE attribution for a manual lead
        assert conn.execute(text('SELECT is_new FROM captured_phones')).scalar() == 0
        assert conn.execute(text('SELECT phones_captured FROM live_metrics')).scalar() == 0
    engine.dispose()
